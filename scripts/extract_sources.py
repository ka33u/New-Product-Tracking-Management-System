from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(paragraph._p.itertext()).strip()


def extract_document(path: Path) -> dict:
    document = Document(path)
    body = []

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = paragraph_text(paragraph)
            if text:
                body.append({"type": "paragraph", "text": text})
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = []
            for row in table.rows:
                rows.append(
                    [
                        "\n".join(
                            filter(
                                None,
                                (paragraph_text(paragraph) for paragraph in cell.paragraphs),
                            )
                        )
                        for cell in row.cells
                    ]
                )
            body.append({"type": "table", "rows": rows})

    headers = []
    footers = []
    for section in document.sections:
        headers.extend(
            text
            for text in (paragraph_text(p) for p in section.header.paragraphs)
            if text
        )
        footers.extend(
            text
            for text in (paragraph_text(p) for p in section.footer.paragraphs)
            if text
        )

    return {
        "file": path.name,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "headers": list(dict.fromkeys(headers)),
        "footers": list(dict.fromkeys(footers)),
        "body": body,
    }


def main() -> None:
    source_dir = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])
    output_dir.mkdir(parents=True, exist_ok=True)

    manifest = []
    for path in sorted(source_dir.glob("*.docx")):
        extracted = extract_document(path)
        output_path = output_dir / f"{path.stem}.json"
        output_path.write_text(
            json.dumps(extracted, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        manifest.append(
            {
                "file": path.name,
                "paragraph_count": extracted["paragraph_count"],
                "table_count": extracted["table_count"],
                "output": str(output_path),
            }
        )

    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
